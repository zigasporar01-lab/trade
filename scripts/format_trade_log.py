"""Turns data/trade_log_<mode>.csv into a formatted, color-coded .xlsx you
can open directly in Excel — no delimiter/regional-settings issues, since
.xlsx isn't a plain-text format.

Usage:
    python scripts/format_trade_log.py                     # paper mode, default paths
    python scripts/format_trade_log.py --mode live
    python scripts/format_trade_log.py --input path/to/log.csv --output path/to/out.xlsx

Safe to re-run anytime — it always rebuilds the workbook fresh from the
current CSV, so it never goes stale relative to your trade history.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")
WIN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
LOSS_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

COLUMNS = [
    ("closed_at", "Closed At", 19, None),
    ("opened_at", "Opened At", 19, None),
    ("mode", "Mode", 8, None),
    ("symbol", "Symbol", 12, None),
    ("token_address", "Token Address", 24, None),
    ("entry_price", "Entry Price ($)", 16, "0.00000000"),
    ("close_price", "Close Price ($)", 16, "0.00000000"),
    ("size_sol", "Size (SOL)", 12, "0.00000"),
    ("pnl_sol", "PnL (SOL)", 12, "+0.00000;-0.00000"),
    ("pnl_pct", "PnL (%)", 10, "+0.0%;-0.0%"),
    ("close_reason", "Close Reason", 14, None),
]


def build_workbook(rows: list[dict]) -> Workbook:
    wb = Workbook()
    trades_ws = wb.active
    trades_ws.title = "Trades"

    for col_idx, (_key, label, width, _fmt) in enumerate(COLUMNS, start=1):
        cell = trades_ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        trades_ws.column_dimensions[get_column_letter(col_idx)].width = width

    numeric_keys = {"entry_price", "close_price", "size_sol", "pnl_sol", "pnl_pct"}
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _label, _width, fmt) in enumerate(COLUMNS, start=1):
            raw = row.get(key, "")
            value = raw
            if key in numeric_keys and raw not in ("", None):
                try:
                    value = float(raw)
                except ValueError:
                    value = raw
            cell = trades_ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            if fmt and isinstance(value, float):
                cell.number_format = fmt

    trades_ws.freeze_panes = "A2"

    last_row = max(len(rows) + 1, 2)
    pnl_col_letter = get_column_letter(next(i for i, c in enumerate(COLUMNS, 1) if c[0] == "pnl_sol"))
    pct_col_letter = get_column_letter(next(i for i, c in enumerate(COLUMNS, 1) if c[0] == "pnl_pct"))
    for col_letter in (pnl_col_letter, pct_col_letter):
        rng = f"{col_letter}2:{col_letter}{last_row}"
        trades_ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=WIN_FILL))
        trades_ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=LOSS_FILL))

    # --- Summary sheet: real formulas referencing Trades, never hardcoded ---
    summary_ws = wb.create_sheet("Summary")
    summary_ws.column_dimensions["A"].width = 22
    summary_ws.column_dimensions["B"].width = 16

    header_a = summary_ws.cell(row=1, column=1, value="Metric")
    header_b = summary_ws.cell(row=1, column=2, value="Value")
    for cell in (header_a, header_b):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    metrics = [
        ("Total Trades", f"=COUNT(Trades!{pnl_col_letter}2:{pnl_col_letter}{last_row})", None),
        ("Wins", f'=COUNTIF(Trades!{pnl_col_letter}2:{pnl_col_letter}{last_row},">0")', None),
        ("Losses", f'=COUNTIF(Trades!{pnl_col_letter}2:{pnl_col_letter}{last_row},"<=0")', None),
        ("Win Rate", "=IF(B2=0,\"n/a\",B3/B2)", "0.0%"),
        ("Total PnL (SOL)", f"=SUM(Trades!{pnl_col_letter}2:{pnl_col_letter}{last_row})", "+0.00000;-0.00000"),
        ("Best Trade (SOL)", f'=IF(B2=0,"n/a",MAX(Trades!{pnl_col_letter}2:{pnl_col_letter}{last_row}))', "+0.00000;-0.00000"),
        ("Worst Trade (SOL)", f'=IF(B2=0,"n/a",MIN(Trades!{pnl_col_letter}2:{pnl_col_letter}{last_row}))', "+0.00000;-0.00000"),
    ]
    for row_idx, (label, formula, fmt) in enumerate(metrics, start=2):
        summary_ws.cell(row=row_idx, column=1, value=label).font = BODY_FONT
        value_cell = summary_ws.cell(row=row_idx, column=2, value=formula)
        value_cell.font = BODY_FONT
        if fmt:
            value_cell.number_format = fmt

    if not rows:
        note = summary_ws.cell(
            row=len(metrics) + 3,
            column=1,
            value="No closed trades yet — this fills in automatically once the bot closes its first position.",
        )
        note.font = Font(name="Arial", italic=True, color="6B7280")

    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="Format a memebot trade log CSV as a colored .xlsx")
    parser.add_argument("--mode", default="paper", help="paper | live (used for default file paths)")
    parser.add_argument("--input", type=Path, default=None, help="Path to the source CSV")
    parser.add_argument("--output", type=Path, default=None, help="Path to write the .xlsx to")
    args = parser.parse_args()

    input_path = args.input or (REPO_ROOT / "data" / f"trade_log_{args.mode}.csv")
    output_path = args.output or input_path.with_suffix(".xlsx")

    if not input_path.exists():
        raise SystemExit(f"No trade log found at {input_path} — has the bot been run yet?")

    with open(input_path, newline="") as f:
        rows = list(csv.DictReader(f))

    wb = build_workbook(rows)
    wb.save(output_path)
    print(f"Wrote {len(rows)} trade(s) to {output_path}")


if __name__ == "__main__":
    main()
